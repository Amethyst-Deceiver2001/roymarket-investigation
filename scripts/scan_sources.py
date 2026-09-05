#!/usr/bin/env python3
"""Run the same entity/product/ИНН/ОГРН extraction id_patterns.py provides
over every file in data/sources/ — the primary-source EGRUL/БФО/press/court-
filing cache — not just Telegram posts. Also doubles as the archiving index:
data/sources/ has no natural directory structure (filenames are a mix of
ФНС UUIDs, ИНН-prefixed EGRUL exports, case numbers, and free-text press
captures) and is only going to grow, so rather than rename files in place
(risky — nothing else should have to track a rename), this script builds a
manifest keyed by extracted ИНН/ОГРН/entity name that acts as the sortable
index. Re-run after adding new files to data/sources/.

Local file processing only (pdftotext, zip/xlsx parsing, html stripping) —
no network calls — same "safe to rerun anytime" category as mine_leads.py.

Per-file extraction (the expensive part — pdftotext/OCR/zip-parsing) is
cached to data/registry/_scan_sources_cache.json, keyed by filename, same
pattern as inventory_sources.py's _source_content_cache.json. A normal run
only (re)processes files not already in the cache, so adding a handful of
new files no longer means re-running OCR/pdftotext over the whole ~500-file
corpus — pass --rescan to force a full reprocess (e.g. after changing the
extraction patterns in id_patterns.py).

Usage:
    python3 scripts/scan_sources.py                  # manifest to stdout
    python3 scripts/scan_sources.py --out data/sources/MANIFEST.md
    python3 scripts/scan_sources.py --leads-out report.md   # also mine for
                                                              # new candidates
    python3 scripts/scan_sources.py --rescan --out data/sources/MANIFEST.md
"""
import argparse
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from extract_ids import scan
from id_patterns import load_known_text, norm

REPO = Path(__file__).resolve().parent.parent
SOURCES = REPO / "data" / "sources"
CACHE = REPO / "data" / "registry" / "_scan_sources_cache.json"


def ocr_pdf_text(path: Path, max_pages: int | None = 5) -> str:
    """OCR fallback for scanned PDFs with no text layer — same pdftoppm +
    tesseract -l rus pattern already established for ФНС УД_*.zip founding
    docs (see egrul_founding_doc_extraction memory), extended here to any
    scanned PDF in data/sources/ (press captures, thank-you letters, etc).
    Only the first `max_pages` pages are OCR'd (default 5, this module's own
    short-document callers keep that default) — enough for the kind of short
    scanned documents this repo actually holds, and keeps a bad/huge PDF from
    hanging the batch scan. A caller expecting longer documents (zakupki
    contract/spec attachments — see mine_zakupki.py's ZAKUPKI_OCR_PAGE_CAP)
    can pass a higher cap; a confirmed real miss (an 11-page scanned spec
    attachment with its item table on pages 6-11, invisible to the default
    5-page cap) is why this became a parameter instead of staying hardcoded."""
    import tempfile
    try:
        with tempfile.TemporaryDirectory() as tmp:
            prefix = str(Path(tmp) / "page")
            cmd = ["pdftoppm", "-r", "300", "-png", str(path), prefix]
            if max_pages is not None:
                cmd[3:3] = ["-l", str(max_pages)]
            subprocess.run(
                cmd, capture_output=True, timeout=max(120, (max_pages or 5) * 10),
            )
            pages = sorted(Path(tmp).glob("page-*.png"))
            chunks = []
            for page in pages:
                r = subprocess.run(
                    ["tesseract", str(page), "-", "-l", "rus"],
                    capture_output=True, timeout=60,
                )
                chunks.append(r.stdout.decode("utf-8", errors="ignore"))
            return "\n".join(chunks)
    except Exception as e:
        return f"[ocr failed: {e}]"


def pdf_text(path: Path, max_pages: int | None = 5) -> str:
    try:
        r = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            capture_output=True, timeout=60,
        )
        text = r.stdout.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[pdftotext failed: {e}]"
    # A real text layer yields more than a few stray characters; anything
    # near-empty on a file that opens fine is a scanned image (2026-08-22:
    # caught on Благодарность.pdf, a thank-you letter scan pdftotext
    # returned zero characters for).
    if len(text.strip()) < 20:
        ocr = ocr_pdf_text(path, max_pages=max_pages)
        if ocr.strip():
            return ocr
    return text


def html_text(path: Path) -> str:
    raw = path.read_bytes().decode("utf-8", errors="ignore")
    try:
        from bs4 import BeautifulSoup
        return BeautifulSoup(raw, "html.parser").get_text(" ")
    except Exception:
        return re.sub(r"<[^>]+>", " ", raw)


def ocr_image_bytes(data: bytes, suffix: str) -> str:
    """OCR a single scanned page — used for TIFF pages found inside ФНС
    УД_*.zip founding-document bundles (see egrul_founding_doc_extraction
    memory). tesseract reads TIFF directly, no pdftoppm conversion step
    needed."""
    import tempfile
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:
            tmp.write(data)
            tmp.flush()
            r = subprocess.run(["tesseract", tmp.name, "-", "-l", "rus"],
                               capture_output=True, timeout=60)
            return r.stdout.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[ocr failed: {e}]"


def pdf_bytes_text(data: bytes, max_pages: int | None = 5) -> str:
    """pdftotext (+ OCR fallback) for a PDF held only as in-memory bytes —
    an inner document inside a zip, not a file on disk."""
    import tempfile
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
            tmp.write(data)
            tmp.flush()
            return pdf_text(Path(tmp.name), max_pages=max_pages)
    except Exception as e:
        return f"[pdf extract failed: {e}]"


def zip_text(path: Path, _depth: int = 0, max_pages: int | None = 5) -> str:
    """Recursively extracts text from a zip's members, including nested
    zips (ФНС УД_*.zip founding-document bundles are a zip containing
    another zip containing scanned TIFF pages or PDFs — see
    egrul_founding_doc_extraction memory). `_depth` guards against a
    pathological zip bomb; founding-doc bundles are at most 2 levels deep
    in practice. `max_pages` threads through to any inner PDF's OCR
    fallback — see ocr_pdf_text's docstring."""
    if _depth > 4:
        return "[zip nesting too deep, stopped]"
    chunks = []
    try:
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                name = info.filename
                try:
                    name = name.encode("cp437").decode("cp866")
                except Exception:
                    pass
                lower = name.lower()
                data = z.read(info.filename)
                if lower.endswith(".xlsx"):
                    try:
                        import io
                        import openpyxl
                        # read_only=True — see mine_zakupki.py's xlsx handler
                        # for why (a large itemized-parts sheet can otherwise
                        # make load_workbook slow enough to look hung).
                        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                        for ws in wb.worksheets:
                            for row in ws.iter_rows(values_only=True):
                                chunks.append(" ".join(str(c) for c in row if c is not None))
                    except Exception as e:
                        chunks.append(f"[xlsx parse failed for {name}: {e}]")
                elif lower.endswith((".txt", ".xml", ".csv", ".md")):
                    chunks.append(data.decode("utf-8", errors="ignore"))
                elif lower.endswith(".pdf"):
                    chunks.append(f"--- {name} ---")
                    chunks.append(pdf_bytes_text(data, max_pages=max_pages))
                elif lower.endswith((".tiff", ".tif", ".jpg", ".jpeg", ".png", ".webp")):
                    chunks.append(f"--- {name} (OCR) ---")
                    chunks.append(ocr_image_bytes(data, Path(name).suffix or ".tiff"))
                elif lower.endswith((".docx", ".doc", ".rtf", ".odt", ".xls")):
                    # 2026-09-05: a real zip member of exactly this shape
                    # (a signed ГК docx sitting next to a .sign file) was
                    # silently invisible before this branch existed — see
                    # office_text/legacy_xls_text's own docstrings for why
                    # this class of file turned out to matter.
                    import tempfile
                    chunks.append(f"--- {name} ---")
                    suffix = Path(name).suffix or ".docx"
                    try:
                        with tempfile.NamedTemporaryFile(suffix=suffix) as tmpf:
                            tmpf.write(data)
                            tmpf.flush()
                            if suffix == ".xls":
                                chunks.append(legacy_xls_text(Path(tmpf.name)))
                            else:
                                chunks.append(office_text(Path(tmpf.name)))
                    except Exception as e:
                        chunks.append(f"[office extract failed for {name}: {e}]")
                elif lower.endswith(".zip"):
                    import io
                    try:
                        with zipfile.ZipFile(io.BytesIO(data)) as _:
                            pass
                        chunks.append(f"--- nested zip {name} ---")
                        chunks.append(_zip_bytes_text(data, _depth + 1, max_pages))
                    except Exception as e:
                        chunks.append(f"[nested zip read failed for {name}: {e}]")
    except Exception as e:
        return f"[zip read failed: {e}]"
    return "\n".join(chunks)


def _zip_bytes_text(data: bytes, depth: int, max_pages: int | None = 5) -> str:
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".zip") as tmp:
        tmp.write(data)
        tmp.flush()
        return zip_text(Path(tmp.name), depth, max_pages)


def office_text(path: Path) -> str:
    """docx/doc/rtf/odt via macOS's built-in `textutil` — no new dependency,
    handles all four formats through one converter. Added 2026-09-05 after
    discovering 212 .docx + 35 .doc + 2 .rtf + 4 .odt files in data/sources/
    had NEVER been text-extracted (extract_text had no branch for them at
    all) — the first .docx checked this way turned out to be a full signed
    ГОЗ contract (ГК «Характер-К», Минобороны РФ) sitting completely unmined."""
    try:
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                           capture_output=True, timeout=60)
        return r.stdout.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[textutil failed: {e}]"


def xlsx_text(path: Path) -> str:
    """Same openpyxl approach zip_text already uses for an xlsx zip member,
    applied to a standalone top-level .xlsx (also unhandled before 2026-09-05)."""
    try:
        import warnings
        import openpyxl
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        chunks = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                chunks.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(chunks)
    except Exception as e:
        return f"[xlsx parse failed: {e}]"


def legacy_xls_text(path: Path) -> str:
    """openpyxl can't read the old binary .xls format at all; LibreOffice
    (already installed as `soffice`) converts it to csv headlessly. Slower
    per-file (~2-4s LibreOffice startup) than the other extractors, but only
    29 files hit this path and it's cached after the first run. First real
    file converted this way was a Мин­обороны РФ государственный контракт
    исполнения report — same pattern as the .docx gap, confirming this was
    worth fixing rather than leaving as a known, accepted gap."""
    import tempfile
    try:
        with tempfile.TemporaryDirectory() as tmp:
            r = subprocess.run(
                ["soffice", "--headless", "--convert-to", "csv",
                 "--outdir", tmp, str(path)],
                capture_output=True, timeout=90,
            )
            outs = list(Path(tmp).glob("*.csv"))
            if not outs:
                return f"[soffice produced no output: {r.stderr.decode(errors='ignore')[:200]}]"
            return outs[0].read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        return f"[soffice convert failed: {e}]"


def image_ocr_text(path: Path) -> str:
    """Top-level scanned image (jpg/png/tif/webp) — reuses ocr_image_bytes,
    already relied on for TIFF pages found inside zip bundles, but never
    called for a standalone image file before 2026-09-05."""
    return ocr_image_bytes(path.read_bytes(), path.suffix)


def rar_text(path: Path) -> str:
    """.rar has no Python stdlib support; `bsdtar` (already on macOS) can
    list and extract the rar3/4-era archives zakupki attachments actually
    use (confirmed 2026-09-05 against a real ZAKUPKI_ATTACH_*.rar — bsdtar
    extracted a nested .docx cleanly). Extracts to a temp dir and reuses
    extract_text on each member, same recursive idea as zip_text's nested-zip
    branch, so a docx/pdf/xlsx inside a rar gets the same handling as one
    inside a zip."""
    import tempfile
    try:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["bsdtar", "-xf", str(path), "-C", tmp],
                           capture_output=True, timeout=60)
            chunks = []
            for member in sorted(Path(tmp).rglob("*")):
                if member.is_file():
                    chunks.append(f"--- {member.relative_to(tmp)} ---")
                    chunks.append(extract_text(member))
            return "\n".join(chunks)
    except Exception as e:
        return f"[rar extract failed: {e}]"


def extract_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf == ".pdf":
        text = pdf_text(path)
    elif suf in (".html", ".mhtml"):
        text = html_text(path)
    elif suf == ".zip":
        text = zip_text(path)
    elif suf in (".csv", ".txt", ".xml", ".md"):
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suf == ".json":
        # scripts/egrul_batch_pull.py writes structured EGRULAPI_*.json
        # pulls here; flattening to a plain string lets the same ИНН/ОГРН/
        # entity-name regex pass over it as everything else.
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suf in (".docx", ".doc", ".rtf", ".odt"):
        text = office_text(path)
    elif suf == ".xlsx":
        text = xlsx_text(path)
    elif suf == ".xls":
        text = legacy_xls_text(path)
    elif suf in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"):
        text = image_ocr_text(path)
    elif suf == ".rar":
        text = rar_text(path)
    else:
        return ""
    # pdftotext -layout pads column gaps with runs of spaces to preserve
    # table alignment — RE_ENTITY_BARE's [ \t]+ separator happily spans
    # those gaps, producing garbage like "АО Российская      Российский"
    # (two unrelated table cells glued into one fake entity name).
    # Collapsing intra-line whitespace runs (not newlines) fixes this.
    return re.sub(r"[ \t]{2,}", " ", text)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=None, help="write manifest here instead of stdout")
    ap.add_argument("--leads-out", type=Path, default=None, help="also write a lead-mining report over data/sources/ text")
    ap.add_argument("--rescan", action="store_true",
                     help="reprocess every file, ignoring the existing per-file cache")
    args = ap.parse_args()

    if not SOURCES.exists():
        raise SystemExit(f"{SOURCES} not found")

    brief_text = load_known_text(REPO)

    cache = {}
    if CACHE.exists() and not args.rescan:
        cache = json.loads(CACHE.read_text(encoding="utf-8"))

    manifest_rows = []
    all_new_entities = {}  # key -> {display, files}
    all_new_people = {}

    # rglob, not iterdir: a manually-unzipped ZAKUPKI_ATTACH folder was found
    # 2026-09-05 sitting one level deep, invisible to a top-level-only scan
    # (2 files, never hashed/typed/mined). cache_key matches bare filename
    # for top-level files (preserves every existing cache hit) and falls
    # back to "subdir/filename" only for nested files, which were never
    # cached under any key before.
    files = sorted(p for p in SOURCES.rglob("*")
                   if p.is_file() and p.name != ".DS_Store")

    def cache_key(p: Path) -> str:
        rel = p.relative_to(SOURCES)
        return str(rel) if len(rel.parts) > 1 else rel.name

    to_process = [p for p in files if cache_key(p) not in cache]
    print(f"{len(files)} files total, {len(cache)} already cached, "
          f"{len(to_process)} to (re)process", file=sys.stderr)

    for path in to_process:
        text = extract_text(path)
        key = cache_key(path)
        if not text.strip():
            cache[key] = {"kind": "?", "inns": [], "ogrns": [], "entities": [], "people": []}
            continue
        hits = scan(text)
        cache[key] = {
            "kind": path.suffix.lstrip("."),
            "inns": hits["ИНН"],
            "ogrns": hits["ОГРН"],
            "entities": hits["Entities (АО/ООО/НПЦ/etc.)"],
            "people": hits["Named individuals (role + name)"],
        }

    CACHE.parent.mkdir(parents=True, exist_ok=True)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    for path in files:
        key = cache_key(path)
        entry = cache.get(key, {"kind": "?", "inns": [], "ogrns": [], "entities": [], "people": []})
        inns = entry["inns"]
        ogrns = entry["ogrns"]
        entities = entry["entities"]
        people = entry["people"]
        manifest_rows.append((key, entry["kind"], inns, ogrns, entities))

        for e in entities:
            ekey = norm(e)
            if not ekey or ekey in brief_text:
                continue
            all_new_entities.setdefault(ekey, {"display": e, "files": set()})["files"].add(key)
        for p in people:
            pkey = norm(p)
            words = [w for w in pkey.split() if len(w) > 1]
            if not words or all(w in brief_text.split() for w in words):
                continue
            all_new_people.setdefault(pkey, {"display": p, "files": set()})["files"].add(key)

    # Manifest: filename -> what it's about, for sorting/lookup without renaming files.
    lines = [f"# data/sources/ manifest — {len(files)} files\n",
             "Auto-generated by `scripts/scan_sources.py`. Files are NOT renamed — this",
             "table is the index. Grep an ИНН/ОГРН/entity name here before re-downloading.\n",
             "| file | type | ИНН | ОГРН | entities mentioned |",
             "|---|---|---|---|---|"]
    for name, kind, inns, ogrns, entities in manifest_rows:
        lines.append(
            f"| `{name}` | {kind} | {', '.join(inns) or '—'} | {', '.join(ogrns) or '—'} | "
            f"{', '.join(entities[:4]) + ('…' if len(entities) > 4 else '') if entities else '—'} |"
        )
    manifest = "\n".join(lines) + "\n"

    if args.out:
        args.out.write_text(manifest, encoding="utf-8")
        print(f"Wrote {args.out}")
    else:
        print(manifest)

    if args.leads_out:
        lead_lines = [f"# New candidates found in data/sources/ (not yet in docs/index.html)\n",
                       "## Entities\n"]
        for key, rec in sorted(all_new_entities.items()):
            lead_lines.append(f"- **{rec['display']}** — {', '.join(sorted(rec['files']))}")
        lead_lines.append("\n## Named individuals\n")
        for key, rec in sorted(all_new_people.items()):
            lead_lines.append(f"- **{rec['display']}** — {', '.join(sorted(rec['files']))}")
        args.leads_out.write_text("\n".join(lead_lines) + "\n", encoding="utf-8")
        print(f"Wrote {args.leads_out}")


if __name__ == "__main__":
    main()
